from __future__ import annotations

import math
import re
import tempfile
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from statistics import median

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image

from .models import PageResult, RenderedPage, VisionElement


FONT_MAP = {
    "sans_serif": "Arial",
    "serif": "Times New Roman",
    "monospace": "Courier New",
    "handwriting": "Segoe Print",
    "unknown": "Arial",
}
ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justified": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
VISUAL_KINDS = {"graphic", "signature", "stamp", "barcode"}
BODY_KINDS = {"paragraph", "list_item", "form_field", "table", "other"}
LIST_MARKER_RE = re.compile(
    r"^\s*(?P<marker>(?:\d+|[A-Za-z])[.)]|[•●▪◦‣⁃*-]|[☐☑☒□■])\s*(?P<text>.*)$"
)
GENERIC_VISUAL_TEXT_RE = re.compile(
    r"^(?:logo|graphic|image|photo|signature|stamp|seal|barcode|qr(?: code)?|"
    r"horizontal line|separator|decorative element)(?:\s+and\s+\w+)?$",
    re.IGNORECASE,
)
LANGUAGE_MAP = {
    "sq": "sq-AL",
    "albanian": "sq-AL",
    "en": "en-US",
    "english": "en-US",
    "fr": "fr-FR",
    "french": "fr-FR",
    "it": "it-IT",
    "italian": "it-IT",
    "de": "de-DE",
    "german": "de-DE",
    "es": "es-ES",
    "spanish": "es-ES",
}


@dataclass(frozen=True, slots=True)
class PageGeometry:
    width_pt: float
    height_pt: float
    left_pt: float
    right_pt: float
    top_pt: float
    bottom_pt: float

    @property
    def usable_width_pt(self) -> float:
        return max(72.0, self.width_pt - self.left_pt - self.right_pt)


@dataclass(slots=True)
class FlowState:
    previous_bottom: float | None = None
    ordered_num_id: int | None = None
    last_outer_number: int | None = None
    restart_numbering: bool = False


@dataclass(frozen=True, slots=True)
class ParsedListItem:
    marker: str
    text: str
    style: str
    level: int
    start: int | None


def export_translation_ready(pages: list[PageResult], output_path: Path) -> Path:
    """Build a semantic, reflowable Word document intended for translation."""
    if not pages:
        raise ValueError("No OCR pages to export.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _set_metadata(document)
    body_size, language = _configure_styles(document, pages)
    numbering = NumberingManager(document)
    state = FlowState()

    with tempfile.TemporaryDirectory(prefix="layoutlens-flow-crops-") as crop_dir_raw:
        crop_dir = Path(crop_dir_raw)
        for page_index, result in enumerate(pages):
            if page_index:
                document.add_section(WD_SECTION.NEW_PAGE)
                state.previous_bottom = None

            section = document.sections[-1]
            geometry = _page_geometry(result)
            _configure_section(section, geometry)
            _build_header_footer(section, result, geometry, crop_dir, language)
            body_elements = _body_elements(result, geometry)
            units = _group_same_line_elements(body_elements)

            for unit in units:
                if isinstance(unit, tuple):
                    paragraph = _add_tabbed_row(
                        document, unit, result.source, geometry, body_size, language
                    )
                    _apply_vertical_rhythm(paragraph, unit[0], result.source, state)
                    state.restart_numbering = False
                    continue

                element = unit
                if element.kind == "table" and element.table_cells:
                    _add_table(document, element, result.source, geometry, body_size, language)
                    state.previous_bottom = element.bbox.y + element.bbox.height
                    state.restart_numbering = True
                elif element.kind in VISUAL_KINDS:
                    _add_visual(
                        document, element, result.source, geometry, crop_dir, language
                    )
                    state.previous_bottom = element.bbox.y + element.bbox.height
                    state.restart_numbering = True
                elif element.kind == "list_item" or _contains_list_marker(element.text):
                    _add_list_element(
                        document,
                        element,
                        result.source,
                        geometry,
                        numbering,
                        state,
                        body_size,
                        language,
                    )
                else:
                    paragraphs = _add_text_element(
                        document, element, result.source, geometry, body_size, language
                    )
                    if paragraphs:
                        _apply_vertical_rhythm(paragraphs[0], element, result.source, state)
                    state.restart_numbering = element.kind in {"title", "heading"}

        document.save(output_path)
    return output_path


def _set_metadata(document: Document) -> None:
    document.core_properties.title = "Translation-ready OCR reconstruction"
    document.core_properties.subject = "Semantic PDF-to-Word OCR"
    document.core_properties.author = "LayoutLens OCR"
    settings = document.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    mode = OxmlElement("w:compatSetting")
    mode.set(qn("w:name"), "compatibilityMode")
    mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    mode.set(qn("w:val"), "15")
    compat.append(mode)
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _configure_styles(document: Document, pages: list[PageResult]) -> tuple[float, str]:
    body_elements = [
        element
        for result in pages
        for element in result.layout.elements
        if element.kind in BODY_KINDS and element.text.strip()
    ]
    weighted_families: Counter[str] = Counter()
    weighted_sizes: list[float] = []
    for element in body_elements:
        weight = max(1, min(120, len(element.text)))
        weighted_families[element.font_family] += weight
        weighted_sizes.extend([element.font_size_pt] * min(weight, 40))

    family_key = weighted_families.most_common(1)[0][0] if weighted_families else "sans_serif"
    family = FONT_MAP[family_key]
    body_size = max(9.0, min(12.0, median(weighted_sizes) if weighted_sizes else 10.5))
    language = _document_language(pages)

    style_tokens = {
        "Normal": (body_size, False, "000000", 0, 4),
        "Title": (max(16.0, min(24.0, body_size * 1.9)), True, "000000", 0, 8),
        "Subtitle": (max(11.0, min(15.0, body_size * 1.2)), False, "000000", 0, 7),
        "Heading 1": (max(14.0, min(19.0, body_size * 1.55)), True, "000000", 10, 4),
        "Heading 2": (max(12.0, min(16.0, body_size * 1.3)), True, "000000", 8, 3),
        "Heading 3": (max(11.0, min(14.0, body_size * 1.15)), True, "000000", 6, 2),
        "List Paragraph": (body_size, False, "000000", 0, 3),
    }
    for name, (size, bold, color, before, after) in style_tokens.items():
        style = document.styles[name]
        style.font.name = family
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), family)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), family)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), family)
        lang = style._element.get_or_add_rPr().find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            style._element.get_or_add_rPr().append(lang)
        lang.set(qn("w:val"), language)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.08
        style_ppr = style._element.get_or_add_pPr()
        inherited_border = style_ppr.find(qn("w:pBdr"))
        if inherited_border is not None:
            style_ppr.remove(inherited_border)

    document.styles["Subtitle"].font.italic = True
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].paragraph_format.keep_with_next = True
    return body_size, language


def _document_language(pages: list[PageResult]) -> str:
    languages = [
        language.strip().lower()
        for result in pages
        for language in result.layout.detected_languages
        if language.strip()
    ]
    if not languages:
        return "en-US"
    common = Counter(languages).most_common(1)[0][0]
    return LANGUAGE_MAP.get(common, common if "-" in common else "en-US")


def _page_geometry(result: PageResult) -> PageGeometry:
    page = result.source
    elements = result.layout.elements
    body = [
        element
        for element in elements
        if element.kind not in {"header", "footer", "page_number"}
        and not _zone_visual(element, "header")
        and not _zone_visual(element, "footer")
        and (element.text.strip() or element.kind in VISUAL_KINDS)
    ]
    if body:
        left_norm = _percentile([element.bbox.x for element in body], 0.12)
        right_norm = 1000.0 - _percentile(
            [element.bbox.x + element.bbox.width for element in body], 0.88
        )
        top_norm = min(element.bbox.y for element in body)
    else:
        left_norm = right_norm = 90.0
        top_norm = 90.0

    footer_items = [
        element
        for element in elements
        if element.kind in {"footer", "page_number"} or _zone_visual(element, "footer")
    ]
    footer_norm = min((element.bbox.y for element in footer_items), default=940.0)
    left = _clamp(page.width_pt * left_norm / 1000.0, 36.0, 90.0)
    right = _clamp(page.width_pt * right_norm / 1000.0, 36.0, 90.0)
    top = _clamp(page.height_pt * top_norm / 1000.0 - 6.0, 54.0, 126.0)
    bottom = _clamp(page.height_pt * (1000.0 - footer_norm) / 1000.0 + 8.0, 42.0, 72.0)
    if left + right > page.width_pt - 144.0:
        left = right = 54.0
    return PageGeometry(page.width_pt, page.height_pt, left, right, top, bottom)


def _configure_section(section, geometry: PageGeometry) -> None:
    section.page_width = Pt(geometry.width_pt)
    section.page_height = Pt(geometry.height_pt)
    section.left_margin = Pt(geometry.left_pt)
    section.right_margin = Pt(geometry.right_pt)
    section.top_margin = Pt(geometry.top_pt)
    section.bottom_margin = Pt(geometry.bottom_pt)
    section.header_distance = Pt(16)
    section.footer_distance = Pt(16)
    section.gutter = Pt(0)


def _build_header_footer(
    section,
    result: PageResult,
    geometry: PageGeometry,
    crop_dir: Path,
    language: str,
) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    headers = [
        element
        for element in result.layout.elements
        if element.kind == "header" or _zone_visual(element, "header")
    ]
    footers = [
        element
        for element in result.layout.elements
        if element.kind in {"footer", "page_number"} or _zone_visual(element, "footer")
    ]
    _fill_story_part(
        section.header,
        headers,
        result.source,
        geometry,
        crop_dir,
        language,
        is_footer=False,
    )
    _fill_story_part(
        section.footer,
        footers,
        result.source,
        geometry,
        crop_dir,
        language,
        is_footer=True,
    )


def _fill_story_part(
    story_part,
    elements: list[VisionElement],
    page: RenderedPage,
    geometry: PageGeometry,
    crop_dir: Path,
    language: str,
    *,
    is_footer: bool,
) -> None:
    paragraph = story_part.paragraphs[0]
    _clear_paragraph(paragraph)
    if not elements:
        return

    visual = [element for element in elements if element.kind in VISUAL_KINDS]
    text = [element for element in elements if element.kind not in VISUAL_KINDS]
    used_first = False
    compact_visuals = [element for element in visual if not _is_rule(element)]
    if compact_visuals:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for index, element in enumerate(sorted(compact_visuals, key=lambda item: item.bbox.x)):
            crop = _crop_region(page, element, crop_dir, f"story-{page.page_number}-{index}")
            width, height = _scaled_picture_size(
                crop,
                max_width=min(geometry.usable_width_pt * 0.32, 150.0),
                max_height=34.0,
            )
            picture = paragraph.add_run().add_picture(
                str(crop), width=Pt(width), height=Pt(height)
            )
            _set_picture_alt(picture, _visual_alt_text(element, page.page_number))
            if index < len(compact_visuals) - 1:
                paragraph.add_run("   ")
        paragraph.paragraph_format.space_after = Pt(2)
        used_first = True

    text_groups = _group_story_text(text)
    for index, group in enumerate(text_groups):
        if len(group) > 1:
            if not used_first and index == 0 and paragraph._p.getparent() is not None:
                paragraph._p.getparent().remove(paragraph._p)
            _add_story_row_table(story_part, group, geometry, language)
            used_first = True
            continue
        target = story_part.add_paragraph() if used_first or index else paragraph
        used_first = True
        element = group[0]
        target.alignment = ALIGN_MAP[element.alignment] if len(group) == 1 else WD_ALIGN_PARAGRAPH.LEFT
        target.paragraph_format.space_before = Pt(0)
        target.paragraph_format.space_after = Pt(1)
        target.paragraph_format.line_spacing = 1.0
        background = next(
            (
                item.background_hex
                for item in group
                if _safe_color(item.background_hex, "FFFFFF") != "FFFFFF"
            ),
            None,
        )
        if background:
            _paragraph_shading(target, background)
        for group_index, item in enumerate(group):
            value = _natural_text(item.text, preserve_lines=True)
            if not value:
                continue
            if group_index:
                target.add_run("\t")
                is_right = item.alignment == "right" or item.bbox.x + item.bbox.width > 820
                tab_anchor = item.bbox.x + item.bbox.width if is_right else item.bbox.x
                tab_position = page.width_pt * tab_anchor / 1000.0 - geometry.left_pt
                alignment = (
                    WD_TAB_ALIGNMENT.RIGHT
                    if is_right
                    else WD_TAB_ALIGNMENT.LEFT
                )
                target.paragraph_format.tab_stops.add_tab_stop(
                    Pt(_clamp(tab_position, 18.0, geometry.usable_width_pt - 4.0)), alignment
                )
            _add_story_content(
                target,
                item,
                value,
                language,
                size=_clamp(item.font_size_pt, 7.0, 11.0),
            )

    if visual and any(_is_rule(element) for element in visual):
        rule_target = story_part.add_paragraph() if used_first else paragraph
        _paragraph_border(rule_target, "top" if is_footer else "bottom")
        rule_target.paragraph_format.space_before = Pt(1)
        rule_target.paragraph_format.space_after = Pt(1)


def _add_story_row_table(
    story_part,
    group: tuple[VisionElement, ...],
    geometry: PageGeometry,
    language: str,
) -> None:
    table = story_part.add_table(rows=1, cols=len(group), width=Pt(geometry.usable_width_pt))
    table.autofit = False
    total_twips = round(geometry.usable_width_pt * 20)
    widths = [total_twips // len(group)] * len(group)
    widths[-1] += total_twips - sum(widths)
    _set_table_geometry(table, total_twips, 0, widths)
    _set_table_no_borders(table)
    _set_table_cell_margins(table, top=20, start=0, bottom=20, end=0)
    background = next(
        (
            item.background_hex
            for item in group
            if _safe_color(item.background_hex, "FFFFFF") != "FFFFFF"
        ),
        None,
    )
    for index, (cell, item) in enumerate(zip(table.rows[0].cells, group)):
        _set_cell_width(cell, widths[index])
        if background:
            _set_cell_shading(cell, background)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
            if index == 0
            else WD_ALIGN_PARAGRAPH.RIGHT
            if index == len(group) - 1
            else WD_ALIGN_PARAGRAPH.CENTER
        )
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        _add_story_content(
            paragraph,
            item,
            _natural_text(item.text, preserve_lines=True),
            language,
            size=_clamp(item.font_size_pt, 7.0, 11.0),
        )


def _add_story_content(
    paragraph,
    element: VisionElement,
    value: str,
    language: str,
    *,
    size: float,
) -> None:
    if element.kind != "page_number":
        run = paragraph.add_run(value)
        _format_run(run, element, language, size=size)
        return

    matches = list(re.finditer(r"\d+", value))
    if not matches:
        run = paragraph.add_run(value)
        _format_run(run, element, language, size=size)
        return

    cursor = 0
    for index, match in enumerate(matches):
        if index >= 2:
            break
        if match.start() > cursor:
            run = paragraph.add_run(value[cursor : match.start()])
            _format_run(run, element, language, size=size)
        _append_word_field(
            paragraph,
            "PAGE" if index == 0 else "NUMPAGES",
            match.group(),
            element,
            language,
            size,
        )
        cursor = match.end()
    if cursor < len(value):
        run = paragraph.add_run(value[cursor:])
        _format_run(run, element, language, size=size)


def _append_word_field(
    paragraph,
    instruction: str,
    placeholder: str,
    element: VisionElement,
    language: str,
    size: float,
) -> None:
    begin = paragraph.add_run()
    begin_node = OxmlElement("w:fldChar")
    begin_node.set(qn("w:fldCharType"), "begin")
    begin._r.append(begin_node)

    code = paragraph.add_run()
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = f" {instruction} "
    code._r.append(instruction_node)

    separate = paragraph.add_run()
    separate_node = OxmlElement("w:fldChar")
    separate_node.set(qn("w:fldCharType"), "separate")
    separate._r.append(separate_node)

    result = paragraph.add_run(placeholder)

    end = paragraph.add_run()
    end_node = OxmlElement("w:fldChar")
    end_node.set(qn("w:fldCharType"), "end")
    end._r.append(end_node)

    for run in (begin, code, separate, result, end):
        _format_run(run, element, language, size=size)


def _group_story_text(elements: list[VisionElement]) -> list[tuple[VisionElement, ...]]:
    ordered = sorted(elements, key=lambda item: (item.bbox.y, item.bbox.x, item.reading_order))
    groups: list[list[VisionElement]] = []
    for element in ordered:
        if not groups or not _story_same_row(groups[-1][0], element):
            groups.append([element])
        else:
            groups[-1].append(element)
    return [tuple(sorted(group, key=lambda item: item.bbox.x)) for group in groups]


def _story_same_row(left: VisionElement, right: VisionElement) -> bool:
    left_top, left_bottom = left.bbox.y, left.bbox.y + left.bbox.height
    right_top, right_bottom = right.bbox.y, right.bbox.y + right.bbox.height
    overlap = min(left_bottom, right_bottom) - max(left_top, right_top)
    minimum_height = max(1.0, min(left.bbox.height, right.bbox.height))
    return overlap / minimum_height >= 0.45 and len(left.text) <= 100 and len(right.text) <= 100


def _body_elements(result: PageResult, geometry: PageGeometry) -> list[VisionElement]:
    del geometry
    return sorted(
        [
            element
            for element in result.layout.elements
            if element.kind not in {"header", "footer", "page_number"}
            and not _zone_visual(element, "header")
            and not _zone_visual(element, "footer")
            and (element.text.strip() or element.kind in VISUAL_KINDS or element.table_cells)
        ],
        key=lambda item: (item.reading_order, item.bbox.y, item.bbox.x),
    )


def _group_same_line_elements(
    elements: list[VisionElement],
) -> list[VisionElement | tuple[VisionElement, ...]]:
    units: list[VisionElement | tuple[VisionElement, ...]] = []
    index = 0
    while index < len(elements):
        current = elements[index]
        if not _row_eligible(current):
            units.append(current)
            index += 1
            continue

        group = [current]
        cursor = index + 1
        while cursor < len(elements) and _same_row(group[-1], elements[cursor]):
            group.append(elements[cursor])
            cursor += 1
        if len(group) > 1:
            units.append(tuple(sorted(group, key=lambda item: item.bbox.x)))
            index = cursor
        else:
            units.append(current)
            index += 1
    return units


def _row_eligible(element: VisionElement) -> bool:
    return (
        element.kind in {"paragraph", "form_field", "other"}
        and 0 < len(element.text.strip()) <= 140
        and len(element.text.splitlines()) <= 2
        and element.bbox.width < 650
    )


def _same_row(left: VisionElement, right: VisionElement) -> bool:
    if not _row_eligible(right):
        return False
    left_top, left_bottom = left.bbox.y, left.bbox.y + left.bbox.height
    right_top, right_bottom = right.bbox.y, right.bbox.y + right.bbox.height
    overlap = min(left_bottom, right_bottom) - max(left_top, right_top)
    minimum_height = max(1.0, min(left.bbox.height, right.bbox.height))
    horizontally_separate = right.bbox.x >= left.bbox.x + left.bbox.width - 5.0
    return overlap / minimum_height >= 0.45 and horizontally_separate


def _add_tabbed_row(
    document: Document,
    elements: tuple[VisionElement, ...],
    page: RenderedPage,
    geometry: PageGeometry,
    body_size: float,
    language: str,
):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(3)
    first_x_pt = page.width_pt * elements[0].bbox.x / 1000.0
    paragraph.paragraph_format.left_indent = Pt(
        _clamp(first_x_pt - geometry.left_pt, 0.0, geometry.usable_width_pt * 0.25)
    )
    base_x = geometry.left_pt + float(paragraph.paragraph_format.left_indent.pt or 0)
    for index, element in enumerate(elements):
        if index:
            paragraph.add_run("\t")
            tab_anchor = (
                element.bbox.x + element.bbox.width
                if element.alignment == "right"
                else element.bbox.x
            )
            target_x = page.width_pt * tab_anchor / 1000.0 - base_x
            alignment = WD_TAB_ALIGNMENT.RIGHT if element.alignment == "right" else WD_TAB_ALIGNMENT.LEFT
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Pt(_clamp(target_x, 18.0, geometry.usable_width_pt - 6.0)), alignment
            )
        run = paragraph.add_run(_natural_text(element.text))
        _format_run(run, element, language, size=_body_run_size(element, body_size))
    return paragraph


def _add_text_element(
    document: Document,
    element: VisionElement,
    page: RenderedPage,
    geometry: PageGeometry,
    body_size: float,
    language: str,
) -> list:
    preserve_lines = element.kind in {"title", "heading"}
    blocks = _natural_blocks(element.text, preserve_lines=preserve_lines)
    paragraphs = []
    for block_index, block in enumerate(blocks):
        if element.kind == "title":
            style = "Title"
        elif element.kind == "heading":
            style = "Subtitle" if element.italic else f"Heading {_heading_level(element, body_size)}"
        else:
            style = "Normal"
        paragraph = document.add_paragraph(style=style)
        paragraph.alignment = ALIGN_MAP[element.alignment]
        paragraph.paragraph_format.widow_control = True
        if block_index:
            paragraph.paragraph_format.space_before = Pt(0)
        if element.kind not in {"title", "heading"}:
            _apply_horizontal_position(paragraph, element, page, geometry)
        run = paragraph.add_run(block)
        size = None if element.kind in {"title", "heading"} else _body_run_size(element, body_size)
        _format_run(run, element, language, size=size)
        if _safe_color(element.background_hex, "FFFFFF") != "FFFFFF":
            _paragraph_shading(paragraph, element.background_hex)
        paragraphs.append(paragraph)
    return paragraphs


def _heading_level(element: VisionElement, body_size: float) -> int:
    if element.font_size_pt >= body_size * 1.28:
        return 1
    if element.font_size_pt >= body_size * 1.12:
        return 2
    return 3


def _apply_horizontal_position(
    paragraph,
    element: VisionElement,
    page: RenderedPage,
    geometry: PageGeometry,
) -> None:
    if element.alignment in {"center", "right"}:
        return
    x_pt = page.width_pt * element.bbox.x / 1000.0
    indent = _clamp(x_pt - geometry.left_pt, 0.0, min(90.0, geometry.usable_width_pt * 0.22))
    if indent >= 4.0:
        paragraph.paragraph_format.left_indent = Pt(indent)


def _apply_vertical_rhythm(
    paragraph,
    element: VisionElement,
    page: RenderedPage,
    state: FlowState,
) -> None:
    if state.previous_bottom is not None:
        gap_norm = max(0.0, element.bbox.y - state.previous_bottom)
        gap_pt = page.height_pt * gap_norm / 1000.0
        current_spacing = paragraph.paragraph_format.space_before
        current = current_spacing.pt if current_spacing is not None else 0.0
        paragraph.paragraph_format.space_before = Pt(max(current, _clamp(gap_pt * 0.35, 0.0, 22.0)))
    state.previous_bottom = element.bbox.y + element.bbox.height


def _add_list_element(
    document: Document,
    element: VisionElement,
    page: RenderedPage,
    geometry: PageGeometry,
    numbering: "NumberingManager",
    state: FlowState,
    body_size: float,
    language: str,
) -> None:
    items = _parse_list_items(element.text, element.kind)
    if not items:
        paragraphs = _add_text_element(document, element, page, geometry, body_size, language)
        if paragraphs:
            _apply_vertical_rhythm(paragraphs[0], element, page, state)
        return

    first_paragraph = None
    for item in items:
        paragraph = document.add_paragraph(style="List Paragraph")
        paragraph.alignment = ALIGN_MAP[element.alignment]
        paragraph.paragraph_format.widow_control = True
        paragraph.paragraph_format.space_after = Pt(2)
        if item.style == "checkbox":
            paragraph.paragraph_format.left_indent = Pt(24 + 18 * item.level)
            paragraph.paragraph_format.first_line_indent = Pt(-18)
            text = f"{item.marker} {item.text}".rstrip()
        else:
            if item.style == "bullet":
                num_id = numbering.bullet_num_id
            else:
                if (
                    state.ordered_num_id is None
                    or (state.restart_numbering and item.start == 1 and item.level == 0)
                ):
                    state.ordered_num_id = numbering.new_ordered_num(item.start or 1)
                num_id = state.ordered_num_id
                if item.level == 0 and item.start is not None:
                    state.last_outer_number = item.start
            numbering.apply(paragraph, num_id, item.level)
            text = item.text
        run = paragraph.add_run(text)
        _format_run(run, element, language, size=_body_run_size(element, body_size))
        if first_paragraph is None:
            first_paragraph = paragraph
    if first_paragraph is not None:
        _apply_vertical_rhythm(first_paragraph, element, page, state)
    state.restart_numbering = False


def _parse_list_items(text: str, kind: str) -> list[ParsedListItem]:
    normalized = _normalize_unicode(text)
    lines = [re.sub(r"\s+", " ", line).strip() for line in normalized.splitlines() if line.strip()]
    raw_items: list[tuple[str, list[str]]] = []
    for line in lines:
        match = LIST_MARKER_RE.match(line)
        if match:
            raw_items.append((match.group("marker"), [match.group("text").strip()]))
        elif raw_items:
            raw_items[-1][1].append(line)
        elif kind == "list_item":
            raw_items.append(("•", [line]))
        else:
            return []

    parsed: list[ParsedListItem] = []
    for marker, parts in raw_items:
        content = _join_wrapped_lines(parts)
        if marker[0].isdigit():
            start = int(re.match(r"\d+", marker).group())
            parsed.append(ParsedListItem(marker, content, "numbered", 0, start))
        elif marker[0].isalpha():
            start = ord(marker[0].lower()) - ord("a") + 1
            parsed.append(ParsedListItem(marker, content, "lettered", 1, start))
        elif marker in {"☐", "☑", "☒", "□", "■"}:
            parsed.append(ParsedListItem(marker, content, "checkbox", 0, None))
        else:
            parsed.append(ParsedListItem(marker, content, "bullet", 0, None))
    return parsed


class NumberingManager:
    def __init__(self, document: Document):
        self.root = document.part.numbering_part.element
        self._next_abstract = self._next_id("w:abstractNum", "w:abstractNumId")
        self._next_num = self._next_id("w:num", "w:numId")
        self._ordered_abstract = self._add_ordered_abstract()
        self._bullet_abstract = self._add_bullet_abstract()
        self.bullet_num_id = self._add_num(self._bullet_abstract)

    def _next_id(self, tag: str, attribute: str) -> int:
        values = [int(node.get(qn(attribute))) for node in self.root.findall(qn(tag))]
        return max(values, default=0) + 1

    def _add_ordered_abstract(self) -> int:
        abstract_id = self._next_abstract
        self._next_abstract += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        formats = (("decimal", "%1."), ("lowerLetter", "%2."), ("lowerRoman", "%3."))
        for level, (number_format, level_text) in enumerate(formats):
            abstract.append(_numbering_level(level, number_format, level_text))
        self.root.append(abstract)
        return abstract_id

    def _add_bullet_abstract(self) -> int:
        abstract_id = self._next_abstract
        self._next_abstract += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level, marker in enumerate(("•", "◦", "▪")):
            abstract.append(_numbering_level(level, "bullet", marker))
        self.root.append(abstract)
        return abstract_id

    def _add_num(self, abstract_id: int, start: int | None = None) -> int:
        num_id = self._next_num
        self._next_num += 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        if start is not None and start != 1:
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), "0")
            start_override = OxmlElement("w:startOverride")
            start_override.set(qn("w:val"), str(start))
            override.append(start_override)
            num.append(override)
        self.root.append(num)
        return num_id

    def new_ordered_num(self, start: int = 1) -> int:
        return self._add_num(self._ordered_abstract, start)

    @staticmethod
    def apply(paragraph, num_id: int, level: int) -> None:
        properties = paragraph._p.get_or_add_pPr()
        old = properties.find(qn("w:numPr"))
        if old is not None:
            properties.remove(old)
        num_properties = OxmlElement("w:numPr")
        level_node = OxmlElement("w:ilvl")
        level_node.set(qn("w:val"), str(max(0, min(2, level))))
        num_node = OxmlElement("w:numId")
        num_node.set(qn("w:val"), str(num_id))
        num_properties.extend((level_node, num_node))
        properties.append(num_properties)


def _numbering_level(level: int, number_format: str, text: str):
    node = OxmlElement("w:lvl")
    node.set(qn("w:ilvl"), str(level))
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    node.append(start)
    if level:
        restart = OxmlElement("w:lvlRestart")
        restart.set(qn("w:val"), "1")
        node.append(restart)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), number_format)
    node.append(fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), text)
    node.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    node.append(justification)
    p_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(720 + level * 540))
    tabs.append(tab)
    p_properties.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(720 + level * 540))
    indent.set(qn("w:hanging"), "360")
    p_properties.append(indent)
    node.append(p_properties)
    if number_format == "bullet":
        r_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        r_properties.append(fonts)
        node.append(r_properties)
    return node


def _add_table(
    document: Document,
    element: VisionElement,
    page: RenderedPage,
    geometry: PageGeometry,
    body_size: float,
    language: str,
) -> None:
    rows = element.table_cells
    columns = max((len(row) for row in rows), default=1)
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_width_pt = min(
        geometry.usable_width_pt,
        max(geometry.usable_width_pt * 0.48, page.width_pt * element.bbox.width / 1000.0),
    )
    indent_pt = _clamp(
        page.width_pt * element.bbox.x / 1000.0 - geometry.left_pt,
        0.0,
        max(0.0, geometry.usable_width_pt - table_width_pt),
    )
    weights = element.table_column_widths
    if len(weights) != columns or sum(weights) <= 0:
        weights = [1.0] * columns
    total_weight = sum(weights)
    total_twips = max(720, columns * 60, round(table_width_pt * 20))
    widths = [max(60, round(total_twips * weight / total_weight)) for weight in weights]
    widths[-1] += total_twips - sum(widths)
    _set_table_geometry(table, total_twips, round(indent_pt * 20), widths)
    _set_table_borders(table, element.table_border_hex)
    _set_table_cell_margins(table, top=90, start=120, bottom=90, end=120)

    table_font_size = _clamp(element.font_size_pt, 7.5, min(11.0, body_size))
    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        is_header = row_index < element.table_header_rows
        if is_header:
            row_properties = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_properties.append(repeat)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = row_values[column_index] if column_index < len(row_values) else ""
            _set_cell_width(cell, widths[column_index])
            if is_header:
                _set_cell_shading(cell, element.table_header_fill_hex)
            elif _safe_color(element.background_hex, "FFFFFF") != "FFFFFF":
                _set_cell_shading(cell, element.background_hex)
            lines = _natural_blocks(value, preserve_lines=True) or [""]
            cell.paragraphs[0].clear()
            for line_index, line in enumerate(lines):
                paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if column_index > 0 and _short_value_column(rows, column_index):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = paragraph.add_run(line)
                _format_run(
                    run,
                    element,
                    language,
                    size=table_font_size,
                    color=(element.table_header_text_hex if is_header else element.color_hex),
                    bold=(True if is_header else element.bold),
                )

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = Pt(1)


def _set_table_geometry(table, total_twips: int, indent_twips: int, widths: list[int]) -> None:
    properties = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        existing = properties.find(qn(tag))
        if existing is not None:
            properties.remove(existing)
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    properties.append(table_width)
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:w"), str(indent_twips))
    table_indent.set(qn("w:type"), "dxa")
    properties.append(table_indent)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)


def _set_table_borders(table, color: str) -> None:
    properties = table._tbl.tblPr
    old = properties.find(qn("w:tblBorders"))
    if old is not None:
        properties.remove(old)
    borders = OxmlElement("w:tblBorders")
    value = _safe_color(color, "808080")
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), value)
        borders.append(border)
    properties.append(borders)


def _set_table_no_borders(table) -> None:
    properties = table._tbl.tblPr
    old = properties.find(qn("w:tblBorders"))
    if old is not None:
        properties.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    properties.append(borders)


def _set_table_cell_margins(table, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def _set_cell_width(cell, width: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:tcW"))
    if node is None:
        node = OxmlElement("w:tcW")
        properties.append(node)
    node.set(qn("w:w"), str(width))
    node.set(qn("w:type"), "dxa")


def _set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), _safe_color(color, "FFFFFF"))


def _short_value_column(rows: list[list[str]], column: int) -> bool:
    values = [row[column] for row in rows if column < len(row)]
    return bool(values) and sum(len(value) for value in values) / len(values) <= 22


def _add_visual(
    document: Document,
    element: VisionElement,
    page: RenderedPage,
    geometry: PageGeometry,
    crop_dir: Path,
    language: str,
) -> None:
    if _is_rule(element):
        paragraph = document.add_paragraph()
        _paragraph_border(paragraph, "bottom")
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(4)
        return
    crop = _crop_region(page, element, crop_dir, f"body-{page.page_number}-{element.reading_order}")
    paragraph = document.add_paragraph()
    center = element.bbox.x + element.bbox.width / 2
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
        if center < 380
        else WD_ALIGN_PARAGRAPH.RIGHT
        if center > 620
        else WD_ALIGN_PARAGRAPH.CENTER
    )
    max_height = 230.0 if element.kind == "graphic" else 185.0
    width, height = _scaled_picture_size(
        crop,
        max_width=min(geometry.usable_width_pt, page.width_pt * element.bbox.width / 1000.0),
        max_height=max_height,
    )
    picture = paragraph.add_run().add_picture(str(crop), width=Pt(width), height=Pt(height))
    _set_picture_alt(picture, _visual_alt_text(element, page.page_number))
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(4)
    visual_text = _natural_text(element.text)
    if visual_text and not GENERIC_VISUAL_TEXT_RE.match(visual_text):
        caption = document.add_paragraph(style="Normal")
        caption.alignment = paragraph.alignment
        run = caption.add_run(visual_text)
        run.italic = True
        run.font.size = Pt(8)
        _set_run_language(run, language)


def _crop_region(
    page: RenderedPage,
    element: VisionElement,
    output_dir: Path,
    name: str,
) -> Path:
    box = element.bbox
    with Image.open(page.image_path) as image:
        left = max(0, math.floor(image.width * box.x / 1000.0))
        top = max(0, math.floor(image.height * box.y / 1000.0))
        right = min(image.width, math.ceil(image.width * (box.x + box.width) / 1000.0))
        bottom = min(image.height, math.ceil(image.height * (box.y + box.height) / 1000.0))
        right = max(left + 1, right)
        bottom = max(top + 1, bottom)
        crop = image.crop((left, top, right, bottom)).convert("RGB")
        output_path = output_dir / f"{name}.png"
        crop.save(output_path, format="PNG", optimize=True)
    return output_path


def _scaled_picture_size(path: Path, *, max_width: float, max_height: float) -> tuple[float, float]:
    with Image.open(path) as image:
        width_px, height_px = image.size
    if width_px <= 0 or height_px <= 0:
        return 1.0, 1.0
    ratio = min(max_width / width_px, max_height / height_px)
    return max(1.0, width_px * ratio), max(1.0, height_px * ratio)


def _format_run(
    run,
    element: VisionElement,
    language: str,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    family = FONT_MAP[element.font_family]
    run.font.name = family
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), family)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), family)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), family)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = element.bold if bold is None else bold
    run.italic = element.italic
    run.underline = element.underline
    run.font.color.rgb = RGBColor.from_string(_safe_color(color or element.color_hex, "000000"))
    _set_run_language(run, language)


def _set_run_language(run, language: str) -> None:
    properties = run._element.get_or_add_rPr()
    lang = properties.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        properties.append(lang)
    lang.set(qn("w:val"), language)


def _set_picture_alt(picture, description: str) -> None:
    doc_properties = picture._inline.docPr
    doc_properties.set("descr", description)
    doc_properties.set("title", description)


def _visual_alt_text(element: VisionElement, page_number: int) -> str:
    value = _natural_text(element.text)
    if value and not GENERIC_VISUAL_TEXT_RE.match(value):
        return value[:240]
    label = {
        "signature": "Signature",
        "stamp": "Official stamp",
        "barcode": "Barcode",
        "graphic": "Document graphic",
    }.get(element.kind, "Document graphic")
    return f"{label} from source page {page_number}"


def _body_run_size(element: VisionElement, body_size: float) -> float:
    if element.kind == "form_field":
        return _clamp(element.font_size_pt, max(8.0, body_size - 1.0), body_size + 1.0)
    return _clamp(element.font_size_pt, max(8.0, body_size - 1.5), body_size + 1.5)


def _natural_blocks(text: str, *, preserve_lines: bool = False) -> list[str]:
    normalized = _normalize_unicode(text)
    raw_blocks = re.split(r"\n\s*\n", normalized)
    blocks = []
    for raw in raw_blocks:
        lines = [line.strip() for line in raw.splitlines() if line.strip()]
        if not lines:
            continue
        value = "\n".join(lines) if preserve_lines else _join_wrapped_lines(lines)
        value = re.sub(r"[ \t]+", " ", value).strip()
        if value:
            blocks.append(value)
    return blocks


def _natural_text(text: str, *, preserve_lines: bool = False) -> str:
    blocks = _natural_blocks(text, preserve_lines=preserve_lines)
    return "\n\n".join(blocks)


def _join_wrapped_lines(lines: list[str]) -> str:
    if not lines:
        return ""
    output = lines[0].strip()
    for raw_next in lines[1:]:
        next_line = raw_next.strip()
        if not next_line:
            continue
        if output.endswith("\u00ad"):
            output = output[:-1] + next_line
        elif output.endswith("-") and next_line[:1].islower():
            output = output[:-1] + next_line
        else:
            output = f"{output} {next_line}"
    return re.sub(r"\s+", " ", output).strip()


def _normalize_unicode(text: str) -> str:
    value = unicodedata.normalize("NFC", text or "")
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = value.replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "")
    return "".join(character for character in value if character in "\n\t" or ord(character) >= 32)


def _contains_list_marker(text: str) -> bool:
    lines = [line for line in _normalize_unicode(text).splitlines() if line.strip()]
    return bool(lines and LIST_MARKER_RE.match(lines[0]))


def _zone_visual(element: VisionElement, zone: str) -> bool:
    if element.kind not in VISUAL_KINDS:
        return False
    if zone == "header":
        return element.bbox.y < 180 and element.bbox.y + element.bbox.height <= 300
    return element.bbox.y >= 900


def _is_rule(element: VisionElement) -> bool:
    return element.kind == "graphic" and element.bbox.width >= 280 and element.bbox.height <= 10


def _paragraph_border(paragraph, edge: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), "808080")
    borders.append(border)


def _paragraph_shading(paragraph, color: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), _safe_color(color, "FFFFFF"))


def _clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _safe_color(value: str, fallback: str) -> str:
    cleaned = (value or "").strip().lstrip("#").upper()
    return cleaned if len(cleaned) == 6 and all(char in "0123456789ABCDEF" for char in cleaned) else fallback


def _percentile(values: list[float], fraction: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    index = int(round((len(ordered) - 1) * fraction))
    return ordered[max(0, min(len(ordered) - 1, index))]


def _clamp(value: float, low: float, high: float) -> float:
    return max(low, min(high, value))
