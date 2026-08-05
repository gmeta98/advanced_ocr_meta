from __future__ import annotations

import json
import math
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Pt
from PIL import Image

from .models import PageResult, RenderedPage, VisionElement
from .wordflow import export_translation_ready


GRAPHIC_KINDS = {"graphic", "signature", "stamp", "barcode"}
FONT_MAP = {
    "sans_serif": "Arial",
    "serif": "Times New Roman",
    "monospace": "Courier New",
    "handwriting": "Segoe Print",
    "unknown": "Arial",
}
ALIGN_MAP = {
    "left": "left",
    "center": "center",
    "right": "right",
    "justified": "both",
}

# python-docx does not register the legacy VML namespaces used by editable
# Word text boxes. Registering them here keeps the generated OOXML namespaced
# and allows Word/LibreOffice to treat each OCR region as an editable shape.
nsmap.setdefault("v", "urn:schemas-microsoft-com:vml")
nsmap.setdefault("o", "urn:schemas-microsoft-com:office:office")


def export_faithful(pages: list[RenderedPage], output_path: Path) -> Path:
    if not pages:
        raise ValueError("No pages to export.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _set_metadata(document, "Exact-look PDF conversion")
    canvas = document.add_paragraph()

    for index, page in enumerate(pages):
        section = document.sections[-1]
        _configure_section(section, page.width_pt, page.height_pt)
        _image_paragraph(canvas)
        run = canvas.add_run()
        picture = run.add_picture(
            str(page.image_path),
            width=Pt(max(1.0, page.width_pt - 2.0)),
            height=Pt(max(1.0, page.height_pt - 2.0)),
        )
        picture._inline.docPr.set("descr", f"Original PDF page {page.page_number}")
        picture._inline.docPr.set("title", f"Original PDF page {page.page_number}")
        if index < len(pages) - 1:
            document.add_section(WD_SECTION.NEW_PAGE)
            _zero_paragraph(document.paragraphs[-1])
            canvas = document.add_paragraph()

    document.save(output_path)
    return output_path


def export_editable(pages: list[PageResult], output_path: Path) -> Path:
    """Export the default translation-ready, naturally reflowing Word file."""
    return export_translation_ready(pages, output_path)


def export_positioned(pages: list[PageResult], output_path: Path) -> Path:
    """Export the legacy page-positioned reconstruction with editable text boxes."""
    if not pages:
        raise ValueError("No OCR pages to export.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _set_metadata(document, "Editable OCR reconstruction")
    shape_id = 1024

    with tempfile.TemporaryDirectory(prefix="layoutlens-crops-") as crop_dir_raw:
        crop_dir = Path(crop_dir_raw)
        canvas = document.add_paragraph()

        for index, result in enumerate(pages):
            page = result.source
            section = document.sections[-1]
            _configure_section(section, page.width_pt, page.height_pt)
            _zero_paragraph(canvas)
            visual_regions = [
                element for element in result.layout.elements if _is_visual_region(element)
            ]

            for element in result.layout.elements:
                if element not in visual_regions and any(
                    _contains(region, element) for region in visual_regions
                ):
                    continue
                shape_id += 1
                x, y, width, height = _box_points(element, page)
                if element in visual_regions:
                    crop = _crop_region(page, element, crop_dir, shape_id)
                    run = canvas.add_run()
                    inline_shape = run.add_picture(
                        str(crop), width=Pt(width), height=Pt(height)
                    )
                    inline_shape._inline.docPr.set(
                        "descr", element.text.strip() or f"Visual from page {page.page_number}"
                    )
                    _inline_to_anchor(
                        inline_shape._inline,
                        x_pt=x,
                        y_pt=y,
                        behind=False,
                        relative_height=shape_id,
                    )
                else:
                    x, y, width, height = _expand_text_box(
                        element, page, x, y, width, height
                    )
                    canvas._p.append(
                        _text_box_run(
                            element,
                            x_pt=x,
                            y_pt=y,
                            width_pt=width,
                            height_pt=height,
                            shape_id=shape_id,
                        )
                    )

            if index < len(pages) - 1:
                document.add_section(WD_SECTION.NEW_PAGE)
                canvas = document.add_paragraph()

        document.save(output_path)
    return output_path


def _is_visual_region(element: VisionElement) -> bool:
    if element.kind in GRAPHIC_KINDS:
        return True
    return (
        element.kind in {"header", "footer"}
        and element.bbox.width >= 800
        and element.bbox.height >= 30
    )


def _contains(region: VisionElement, element: VisionElement) -> bool:
    outer = region.bbox
    inner = element.bbox
    tolerance = 6.0
    return (
        inner.x >= outer.x - tolerance
        and inner.y >= outer.y - tolerance
        and inner.x + inner.width <= outer.x + outer.width + tolerance
        and inner.y + inner.height <= outer.y + outer.height + tolerance
    )


def export_layout_json(pages: list[PageResult], output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    data = {
        "coordinate_space": "normalized_0_1000",
        "pages": [
            {
                "source": {
                    "page_number": result.source.page_number,
                    "width_pt": result.source.width_pt,
                    "height_pt": result.source.height_pt,
                },
                "layout": result.layout.model_dump(mode="json"),
            }
            for result in pages
        ],
    }
    output_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return output_path


def _configure_section(section, width_pt: float, height_pt: float) -> None:
    section.page_width = Pt(width_pt)
    section.page_height = Pt(height_pt)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)
    section.gutter = Pt(0)


def _set_metadata(document: Document, title: str) -> None:
    document.core_properties.title = title
    document.core_properties.subject = "Layout-preserving OCR"
    document.core_properties.author = "LayoutLens OCR"
    settings = document.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    mode = OxmlElement("w:compatSetting")
    mode.set(qn("w:name"), "compatibilityMode")
    mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    mode.set(qn("w:val"), "15")
    compat.append(mode)


def _zero_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    paragraph.paragraph_format.keep_with_next = False


def _image_paragraph(paragraph) -> None:
    _zero_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing = paragraph._p.get_or_add_pPr().find(qn("w:spacing"))
    if spacing is not None:
        spacing.attrib.pop(qn("w:line"), None)
        spacing.attrib.pop(qn("w:lineRule"), None)


def _inline_to_anchor(
    inline,
    *,
    x_pt: float,
    y_pt: float,
    behind: bool,
    relative_height: int = 0,
) -> None:
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": str(max(0, relative_height)),
        "behindDoc": "1" if behind else "0",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(name, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "page")
    offset_h = OxmlElement("wp:posOffset")
    offset_h.text = str(round(x_pt * 12700))
    position_h.append(offset_h)
    anchor.append(position_h)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "page")
    offset_v = OxmlElement("wp:posOffset")
    offset_v.text = str(round(y_pt * 12700))
    position_v.append(offset_v)
    anchor.append(position_v)

    for tag in ("wp:extent", "wp:effectExtent"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(deepcopy(child))

    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(deepcopy(child))

    inline.getparent().replace(inline, anchor)


def _box_points(element: VisionElement, page: RenderedPage) -> tuple[float, float, float, float]:
    box = element.bbox
    x = page.width_pt * box.x / 1000.0
    y = page.height_pt * box.y / 1000.0
    width = max(2.0, page.width_pt * box.width / 1000.0)
    height = max(2.0, page.height_pt * box.height / 1000.0)
    return x, y, min(width, page.width_pt - x), min(height, page.height_pt - y)


def _crop_region(
    page: RenderedPage,
    element: VisionElement,
    output_dir: Path,
    shape_id: int,
) -> Path:
    box = element.bbox
    with Image.open(page.image_path) as image:
        left = max(0, math.floor(image.width * box.x / 1000.0))
        top = max(0, math.floor(image.height * box.y / 1000.0))
        right = min(image.width, math.ceil(image.width * (box.x + box.width) / 1000.0))
        bottom = min(image.height, math.ceil(image.height * (box.y + box.height) / 1000.0))
        if right <= left or bottom <= top:
            right = min(image.width, left + 1)
            bottom = min(image.height, top + 1)
        crop = image.crop((left, top, right, bottom)).convert("RGB")
        output_path = output_dir / f"crop-{page.page_number:04d}-{shape_id}.png"
        crop.save(output_path, format="PNG", optimize=True)
    return output_path


def _expand_text_box(
    element: VisionElement,
    page: RenderedPage,
    x: float,
    y: float,
    width: float,
    height: float,
) -> tuple[float, float, float, float]:
    if element.kind == "table":
        return x, y, width, min(page.height_pt - y, height + 2.0)
    lines = max(1, len(element.text.splitlines()))
    line_factor = 1.85 if element.kind in {"form_field", "footer", "page_number"} else 1.45
    needed = element.font_size_pt * lines * line_factor + 3.0
    expansion = max(0.0, needed - height)
    y = max(0.0, y - min(1.5, expansion * 0.2))
    height = min(page.height_pt - y, max(height + 2.0, needed))
    longest_line = max((len(line) for line in element.text.splitlines()), default=0)
    estimated_width = longest_line * element.font_size_pt * 0.52 + 3.0
    width = min(page.width_pt - x, max(width, estimated_width))
    if element.kind in {"form_field", "footer", "page_number"}:
        y = max(0.0, y - 2.0)
    return x, y, width, height


def _text_box_run(
    element: VisionElement,
    *,
    x_pt: float,
    y_pt: float,
    width_pt: float,
    height_pt: float,
    shape_id: int,
):
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    no_proof = OxmlElement("w:noProof")
    run_properties.append(no_proof)
    run.append(run_properties)

    pict = OxmlElement("w:pict")
    shape = OxmlElement("v:shape")
    shape.set("id", f"LayoutLensBox{shape_id}")
    shape.set("type", "#_x0000_t202")
    shape.set("style", _shape_style(x_pt, y_pt, width_pt, height_pt, shape_id, text=True))
    shape.set("filled", "f")
    shape.set("stroked", "f")
    shape.set(qn("o:allowincell"), "f")

    textbox = OxmlElement("v:textbox")
    textbox.set("inset", "0,0,0,0")
    content = OxmlElement("w:txbxContent")
    if element.kind == "table" and element.table_cells:
        content.append(_table_xml(element, width_pt))
    else:
        _append_text_paragraphs(content, element, width_pt, height_pt)
    textbox.append(content)
    shape.append(textbox)
    pict.append(shape)
    run.append(pict)
    return run


def _shape_style(
    x_pt: float,
    y_pt: float,
    width_pt: float,
    height_pt: float,
    shape_id: int,
    *,
    text: bool = False,
) -> str:
    values = [
        "position:absolute",
        f"margin-left:{x_pt:.2f}pt",
        f"margin-top:{y_pt:.2f}pt",
        f"width:{width_pt:.2f}pt",
        f"height:{height_pt:.2f}pt",
        f"z-index:{shape_id}",
        "mso-position-horizontal-relative:page",
        "mso-position-vertical-relative:page",
        "mso-wrap-style:none",
    ]
    if text:
        values.append("v-text-anchor:top")
    return ";".join(values)


def _append_text_paragraphs(
    content, element: VisionElement, width_pt: float, height_pt: float
) -> None:
    lines = element.text.splitlines() or [""]
    line_count = max(1, len(lines))
    longest_line = max((len(line) for line in lines), default=1)
    horizontal_fit = max(5.0, width_pt / max(1.0, longest_line * 0.52))
    size_pt = min(
        element.font_size_pt,
        max(5.0, height_pt / (line_count * 1.1)),
        horizontal_fit,
    )
    for line in lines:
        paragraph = OxmlElement("w:p")
        properties = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), str(max(100, round(size_pt * 20 * 1.05))))
        spacing.set(qn("w:lineRule"), "atLeast")
        properties.append(spacing)
        justify = OxmlElement("w:jc")
        justify.set(qn("w:val"), ALIGN_MAP[element.alignment])
        properties.append(justify)
        paragraph.append(properties)
        paragraph.append(_text_run_xml(line, element, size_pt))
        content.append(paragraph)


def _text_run_xml(
    text: str,
    element: VisionElement,
    size_pt: float,
    *,
    color_override: str | None = None,
    bold_override: bool | None = None,
):
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    family = FONT_MAP[element.font_family]
    fonts.set(qn("w:ascii"), family)
    fonts.set(qn("w:hAnsi"), family)
    fonts.set(qn("w:eastAsia"), family)
    properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(max(10, round(size_pt * 2))))
    properties.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(max(10, round(size_pt * 2))))
    properties.append(size_cs)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _safe_color(color_override or element.color_hex))
    properties.append(color)
    if element.bold if bold_override is None else bold_override:
        properties.append(OxmlElement("w:b"))
    if element.italic:
        properties.append(OxmlElement("w:i"))
    if element.underline:
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.append(underline)
    run.append(properties)
    parts = text.split("\n")
    for index, part in enumerate(parts):
        text_node = OxmlElement("w:t")
        text_node.set(qn("xml:space"), "preserve")
        text_node.text = part
        run.append(text_node)
        if index < len(parts) - 1:
            run.append(OxmlElement("w:br"))
    return run


def _table_xml(element: VisionElement, width_pt: float):
    rows = element.table_cells
    columns = max((len(row) for row in rows), default=1)
    width_twips = max(120, round(width_pt * 20))
    weights = element.table_column_widths
    if len(weights) != columns or sum(weights) <= 0:
        weights = [1.0] * columns
    total_weight = sum(weights)
    column_widths = [max(60, round(width_twips * weight / total_weight)) for weight in weights]

    table = OxmlElement("w:tbl")
    properties = OxmlElement("w:tblPr")
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:w"), str(width_twips))
    table_width.set(qn("w:type"), "dxa")
    properties.append(table_width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    borders = OxmlElement("w:tblBorders")
    border_color = _safe_color(element.table_border_hex)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), border_color)
        borders.append(border)
    properties.append(borders)
    table.append(properties)

    grid = OxmlElement("w:tblGrid")
    for column_twips in column_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(column_twips))
        grid.append(grid_col)
    table.append(grid)

    size_pt = max(5.0, min(element.font_size_pt, 12.0))
    for row_index, row_values in enumerate(rows):
        row = OxmlElement("w:tr")
        is_header = row_index < element.table_header_rows
        for column_index in range(columns):
            cell = OxmlElement("w:tc")
            cell_properties = OxmlElement("w:tcPr")
            cell_width = OxmlElement("w:tcW")
            cell_width.set(qn("w:w"), str(column_widths[column_index]))
            cell_width.set(qn("w:type"), "dxa")
            cell_properties.append(cell_width)
            if is_header:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), _safe_color(element.table_header_fill_hex))
                cell_properties.append(shading)
            elif _safe_color(element.background_hex) != "FFFFFF":
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), _safe_color(element.background_hex))
                cell_properties.append(shading)
            cell.append(cell_properties)
            value = row_values[column_index] if column_index < len(row_values) else ""
            value_lines = value.splitlines() or [""]
            for line_index, line in enumerate(value_lines):
                paragraph = OxmlElement("w:p")
                p_properties = OxmlElement("w:pPr")
                spacing = OxmlElement("w:spacing")
                spacing.set(qn("w:before"), "0")
                spacing.set(qn("w:after"), "0")
                p_properties.append(spacing)
                paragraph.append(p_properties)
                paragraph.append(
                    _text_run_xml(
                        line,
                        element,
                        size_pt,
                        color_override=(
                            element.table_header_text_hex if is_header else element.color_hex
                        ),
                        bold_override=(
                            True
                            if is_header or (len(value_lines) > 1 and line_index == 0)
                            else None
                        ),
                    )
                )
                cell.append(paragraph)
            row.append(cell)
        table.append(row)
    return table


def _safe_color(value: str) -> str:
    cleaned = value.strip().lstrip("#").upper()
    return cleaned if len(cleaned) == 6 and all(c in "0123456789ABCDEF" for c in cleaned) else "000000"
